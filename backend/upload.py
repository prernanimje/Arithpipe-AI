import csv
import io
import json
import traceback
from fastapi import APIRouter, UploadFile, File, HTTPException
from app.db.database import get_connection
from openpyxl import load_workbook

router = APIRouter()

@router.post("/upload")
async def upload_file(file: UploadFile = File(...)):
    try:
        # ---------- BASIC VALIDATION ----------
        if not file or not file.filename:
            raise HTTPException(status_code=400, detail="No file selected")

        content = await file.read()
        filename = file.filename.lower()

        if not content:
            raise HTTPException(status_code=400, detail="Uploaded file is empty")

        print("📁 Uploaded file:", filename)
        print("📦 File size:", len(content))

        data_rows = []

        # ---------- CSV ----------
        if filename.endswith(".csv"):
            try:
                decoded = content.decode("utf-8-sig")
            except UnicodeDecodeError:
                raise HTTPException(status_code=400, detail="CSV must be UTF-8 encoded")

            reader = csv.DictReader(io.StringIO(decoded))
            headers = [h.lower().strip() for h in (reader.fieldnames or [])]

            print("📄 CSV headers:", headers)

            if "label" not in headers or "value" not in headers:
                raise HTTPException(
                    status_code=400,
                    detail="CSV must contain 'label' and 'value' columns"
                )

            for row in reader:
                label = row.get("label")
                value = row.get("value")

                if label is None or value is None:
                    continue

                data_rows.append({
                    "label": str(label).strip(),
                    "value": float(value)
                })

        # ---------- JSON ----------
        elif filename.endswith(".json"):
            try:
                json_data = json.loads(content)
            except json.JSONDecodeError:
                raise HTTPException(status_code=400, detail="Invalid JSON format")

            if not isinstance(json_data, list):
                raise HTTPException(status_code=400, detail="JSON must be a list of objects")

            for item in json_data:
                if not isinstance(item, dict):
                    continue

                label = item.get("label")
                value = item.get("value")

                if label is None or value is None:
                    continue

                data_rows.append({
                    "label": str(label).strip(),
                    "value": float(value)
                })

        # ---------- EXCEL ----------
        elif filename.endswith(".xlsx"):
            try:
                wb = load_workbook(io.BytesIO(content), data_only=True)
            except Exception:
                raise HTTPException(status_code=400, detail="Invalid Excel file")

            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))

            print("📊 Total Excel rows:", len(rows))

            if len(rows) < 2:
                raise HTTPException(status_code=400, detail="Excel file has no data")

            # Normalize headers
            headers = [
                str(h).strip().lower()
                for h in rows[0]
                if h is not None
            ]

            print("🧾 Excel headers:", headers)

            if "label" not in headers or "value" not in headers:
                raise HTTPException(
                    status_code=400,
                    detail="Excel must contain 'label' and 'value' columns"
                )

            label_idx = headers.index("label")
            value_idx = headers.index("value")

            for row in rows[1:]:
                if not row:
                    continue

                if len(row) <= max(label_idx, value_idx):
                    continue

                label = row[label_idx]
                value = row[value_idx]

                if label is None or value is None:
                    continue

                data_rows.append({
                    "label": str(label).strip(),
                    "value": float(value)
                })

        else:
            raise HTTPException(status_code=400, detail="Unsupported file type")

        # ---------- FINAL VALIDATION ----------
        if not data_rows:
            raise HTTPException(status_code=400, detail="No valid data found in file")

        # ---------- SAVE TO DATABASE ----------
        conn = get_connection()
        cur = conn.cursor()

        cur.execute(
            "INSERT INTO uploads (filename, file_data) VALUES (%s, %s) RETURNING id",
            (file.filename, content)
        )

        file_id = cur.fetchone()[0]
        conn.commit()
        cur.close()
        conn.close()

        print("✅ Upload successful. Rows:", len(data_rows))

        return {
            "file_id": file_id,
            "data": data_rows
        }

    except HTTPException as e:
        print("❌ HTTP ERROR:", e.detail)
        raise e

    except Exception as e:
        print("🔥 INTERNAL SERVER ERROR")
        traceback.print_exc()
        raise HTTPException(status_code=500, detail="Internal server error")
